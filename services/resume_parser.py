"""
Resume Parser Service
Extracts text from PDF and DOCX files safely and cleanly.
"""

import os
import logging
from pathlib import Path
from typing import Tuple

logger = logging.getLogger("resume_parser")

def extract_text_from_pdf(filepath: str) -> str:
    """Extract full raw text content from PDF resume."""
    from pypdf import PdfReader
    extracted_text = []

    try:
        reader = PdfReader(filepath)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_text.append(page_text)
        
        full_text = "\n".join(extracted_text).strip()
        if not full_text:
            logger.warning(f"PDF extraction yielded empty text for: {filepath}")
        return full_text
    except Exception as e:
        logger.error(f"Error extracting text from PDF {filepath}: {e}")
        raise ValueError(f"Failed to read PDF file: {str(e)}")

def extract_text_from_docx(filepath: str) -> str:
    """Extract full text content from DOCX resume."""
    from docx import Document
    extracted_text = []

    try:
        doc = Document(filepath)
        for para in doc.paragraphs:
            if para.text.strip():
                extracted_text.append(para.text.strip())
        
        # Also extract table text if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    extracted_text.append(row_text)

        full_text = "\n".join(extracted_text).strip()
        return full_text
    except Exception as e:
        logger.error(f"Error extracting text from DOCX {filepath}: {e}")
        raise ValueError(f"Failed to read DOCX file: {str(e)}")

def extract_text_from_file(filepath: str) -> Tuple[str, str]:
    """
    Detects file extension and extracts text.
    Returns: (extracted_text, file_type)
    """
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath), "pdf"
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath), "docx"
    else:
        raise ValueError(f"Unsupported resume format: {ext}")
